#!/usr/bin/env python3
"""
ورود «فهرست عناوین مصوبات» سامانه ملی قوانین و مقررات (qavanin.ir) به مخزن

خروجی: data/raw/qavanin-index/qavanin-list.tsv.gz  (UTF-8، gzip قطعی/تکرارپذیر)
ستون‌ها: id <TAB> title <TAB> approval_date <TAB> approval_authority

ورودی‌های پشتیبانی‌شده:
  1) فایل Excel خزنده متن‌باز abdal (ستون‌های pid, name, date, approval) — همان منبع نسخه فعلی:
       python scripts/pipeline/import_qavanin_list.py --xlsx data.xlsx
  2) خروجی TSV فرمان `list` ابزار برداشت همین مخزن (scripts/scraper/qavanin_playwright.py list):
       python scripts/pipeline/import_qavanin_list.py --tsv scripts/scraper/output/qavanin-list.tsv
     با --merge ردیف‌های جدید به فهرست موجود افزوده و ردیف‌های هم‌شناسه به‌روز می‌شوند.

اصل امانت: متن عنوان، تاریخ و مرجع تصویب «همان‌گونه که در منبع است» ذخیره می‌شود؛ تنها کاراکترهای
TAB/CR/LF درون فیلدها به فاصله تبدیل و فاصله‌های ابتدا/انتها حذف می‌شود. نرمال‌سازی نگارشی (ی/ک عربی،
نیم‌فاصله و …) فقط هنگام ساخت بسته‌های اپ (scripts/bundle-data.ts) و بدون تغییر این فایل انجام می‌شود.
"""

from __future__ import annotations

import argparse
import csv
import gzip
import io
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "data" / "raw" / "qavanin-index" / "qavanin-list.tsv.gz"
HEADER = ["id", "title", "approval_date", "approval_authority"]
_WS_RE = re.compile(r"[\t\r\n]+")


def clean(value: object) -> str:
    if value is None:
        return ""
    s = str(value)
    if s in ("None", "nan"):
        return ""
    return _WS_RE.sub(" ", s).strip()


def read_xlsx(path: Path) -> list[list[str]]:
    try:
        import openpyxl  # type: ignore
    except ImportError:  # pragma: no cover
        sys.exit("openpyxl لازم است: pip install -r scripts/pipeline/requirements.txt")
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    header = [str(h) for h in next(rows)]
    idx = {name: header.index(name) for name in ("pid", "name", "date", "approval")}
    out: list[list[str]] = []
    for r in rows:
        pid = clean(r[idx["pid"]])
        if not pid:
            continue
        out.append([str(int(float(pid))), clean(r[idx["name"]]), clean(r[idx["date"]]), clean(r[idx["approval"]])])
    return out


def read_tsv(path: Path) -> list[list[str]]:
    opener = gzip.open if path.suffix == ".gz" else open
    with opener(path, "rt", encoding="utf-8", newline="") as f:  # type: ignore[operator]
        reader = csv.reader(f, delimiter="\t", quoting=csv.QUOTE_NONE)
        header = next(reader)
        if header[:4] != HEADER:
            sys.exit(f"سرستون نامعتبر در {path}: {header}")
        return [[clean(c) for c in row[:4]] for row in reader if row and row[0].strip()]


def write_tsv(rows: list[list[str]], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    buf = io.StringIO()
    buf.write("\t".join(HEADER) + "\n")
    for row in rows:
        buf.write("\t".join(row) + "\n")
    data = buf.getvalue().encode("utf-8")
    # mtime=0 و بدون نام فایل → خروجی بایت‌به‌بایت تکرارپذیر
    with open(path, "wb") as raw:
        with gzip.GzipFile(filename="", mode="wb", fileobj=raw, compresslevel=9, mtime=0) as gz:
            gz.write(data)


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    src = ap.add_mutually_exclusive_group(required=True)
    src.add_argument("--xlsx", type=Path, help="Excel خزنده abdal (pid, name, date, approval)")
    src.add_argument("--tsv", type=Path, help="خروجی TSV فرمان list ابزار برداشت")
    ap.add_argument("--merge", action="store_true", help="ادغام با فهرست موجود به‌جای جایگزینی")
    ap.add_argument("--out", type=Path, default=OUT)
    args = ap.parse_args()

    rows = read_xlsx(args.xlsx) if args.xlsx else read_tsv(args.tsv)
    by_id: dict[str, list[str]] = {}
    if args.merge and args.out.exists():
        for row in read_tsv(args.out):
            by_id[row[0]] = row
    before = len(by_id)
    for row in rows:
        by_id[row[0]] = row
    merged = sorted(by_id.values(), key=lambda r: int(r[0]))
    write_tsv(merged, args.out)
    print(f"✓ {len(merged):,} عنوان → {args.out.relative_to(ROOT)} (+{len(merged) - before:,} جدید)")


if __name__ == "__main__":
    main()
